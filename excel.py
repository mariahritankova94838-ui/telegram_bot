import openpyxl
import os
from datetime import datetime, timedelta, timezone

FILE = "students.xlsx"


def get_book():
    if not os.path.exists(FILE):
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.append([
            "ФИО", "Телефон", "Telegram",
            "Класс", "Предмет", "Дата",
            "Время", "Формат"
        ])
        wb.save(FILE)
    return openpyxl.load_workbook(FILE)


def add_student(data):
    wb = get_book()
    sheet = wb.active
    
   
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[2] == data[2] and row[5] == data[5] and row[6] == data[6]:
            return False  
    sheet.append(data)
    wb.save(FILE)
    return True


def get_all():
    wb = get_book()
    sheet = wb.active
    return list(sheet.iter_rows(min_row=2, values_only=True))


def clear():
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append([
        "ФИО", "Телефон", "Telegram",
        "Класс", "Предмет", "Дата",
        "Время", "Формат"
    ])
    wb.save(FILE)


def get_today():
    tz_moscow = timezone(timedelta(hours=3))
    today = datetime.now(tz_moscow).strftime("%d.%m")

    rows = get_all()
    today_lessons = [r for r in rows if r[5] == today]
    return sorted(today_lessons, key=lambda x: x[6] if x[6] else "")
